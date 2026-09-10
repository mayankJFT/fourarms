"""Document ingestion pipeline: extract → chunk → embed → upsert to Pinecone."""

from __future__ import annotations

import logging
import uuid
from datetime import datetime, timezone
from typing import List, Optional

from app.config import settings
from app.documents.chunker import chunk_text
from app.documents.embeddings import embed_texts
from app.documents.models import Chunk, Document
from app.documents.ocr import is_scanned_pdf, ocr_pdf_pages
from app.documents.office_convert import convert_legacy_office_doc, is_legacy_office_format
from app.documents.pinecone_client import pinecone_client

logger = logging.getLogger(__name__)


# ── Table helpers ──────────────────────────────────────────────────────────────

def _table_to_html(rows: List[List[Optional[str]]]) -> str:
    """Render a table (list of rows, each a list of cell values) as an HTML
    <table>. Keeping the whole table as one HTML block — instead of flattening
    it into tab/space-separated prose and letting the word-count chunker slice
    it into arbitrary pieces — preserves row/column structure for the LLM and
    keeps every row together in a single retrievable chunk."""
    def esc(v: Optional[str]) -> str:
        s = "" if v is None else str(v)
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    lines = ["<table>"]
    for row in rows:
        cells = "".join(f"<td>{esc(c)}</td>" for c in row)
        lines.append(f"<tr>{cells}</tr>")
    lines.append("</table>")
    return "\n".join(lines)


def _is_meaningful_table(rows: List[List[Optional[str]]]) -> bool:
    """Skip tables that are empty or just a single stray cell (extraction noise)."""
    non_empty_cells = sum(1 for row in rows for c in row if c and str(c).strip())
    return len(rows) >= 2 and non_empty_cells >= 3


# ── Text extractors ───────────────────────────────────────────────────────────

def _extract_pdf(file_path: str, document: Document) -> List[dict]:
    """Return list of {text, page_number, section_title} dicts from a PDF."""
    pages_data = []

    scanned = is_scanned_pdf(file_path)
    document.is_scanned = scanned

    if scanned:
        ocr_results = ocr_pdf_pages(file_path)
        confidences = []
        for page_num, text, conf in ocr_results:
            pages_data.append({"text": text, "page_number": page_num, "section_title": None})
            if conf > 0:
                confidences.append(conf)
        document.ocr_confidence = (
            sum(confidences) / len(confidences) if confidences else 0.0
        )
        document.page_count = len(ocr_results)
    else:
        import pdfplumber
        from app.documents.chunker import _clean_page_text, _split_into_sections

        # Heading detection used to happen per-page, independently, down in
        # chunk_text() — so a section that starts on page 3 and continues onto
        # page 4 lost its heading entirely for the page-4 continuation (nothing
        # on page 4 itself looks like a heading), leaving those chunks with
        # section_title=None and no way for retrieval or the LLM reranker to
        # know what they're about. Detecting headings here instead, threading
        # the current section across the page loop (like _extract_docx already
        # does for Word files), keeps a section's identity continuous across
        # page breaks — pages_data is now pre-labelled, so chunk_text() below
        # just windows within each already-known section instead of guessing.
        current_section: Optional[str] = None
        with pdfplumber.open(file_path) as pdf:
            document.page_count = len(pdf.pages)
            for page_num, page in enumerate(pdf.pages, start=1):
                raw_text = page.extract_text() or ""
                cleaned = _clean_page_text(raw_text)
                for heading, body in _split_into_sections(cleaned):
                    if heading:
                        current_section = heading
                    if body.strip():
                        pages_data.append({
                            "text": body,
                            "page_number": page_num,
                            "section_title": current_section,
                        })

                # Tables extracted as plain text get flattened into jumbled prose by
                # extract_text() above (columns run together) and, worse, get sliced
                # across multiple chunk boundaries by the word-count chunker — losing
                # row/column structure right when it matters most (specs, deliverables,
                # pricing tables). Extract each table separately as one atomic HTML chunk.
                #
                # find_tables() (unlike extract_tables()) also hands back each table's
                # bounding box, which lets us grab the line of text immediately above
                # it as a pseudo-heading (e.g. "Pricing & Timeline Estimates") and
                # prefix it into the embedded text — a bare "<table><tr><td>..." embeds
                # nowhere near a query like "costing", but "Pricing & Timeline
                # Estimates:\n<table>..." does.
                try:
                    table_objs = page.find_tables()
                except Exception as exc:
                    logger.warning("Table extraction failed on page %d: %s", page_num, exc)
                    table_objs = []
                for t_idx, table_obj in enumerate(table_objs, start=1):
                    try:
                        rows = table_obj.extract()
                    except Exception:
                        continue
                    if not _is_meaningful_table(rows):
                        continue
                    heading = None
                    try:
                        top = table_obj.bbox[1]
                        if top > 0:
                            above_text = page.crop((0, 0, page.width, top)).extract_text() or ""
                            lines = [ln.strip() for ln in above_text.splitlines() if ln.strip()]
                            if lines:
                                heading = lines[-1][:120]
                    except Exception:
                        heading = None
                    html = _table_to_html(rows)
                    pages_data.append({
                        "text": f"{heading}:\n{html}" if heading else html,
                        "page_number": page_num,
                        "section_title": heading or f"Table {t_idx} (page {page_num})",
                        "is_table": True,
                    })

    return pages_data


def _iter_block_items(doc):
    """Yield each paragraph and table of a docx Document in true document order.

    python-docx exposes doc.paragraphs and doc.tables as two separate flat lists
    with no positional relationship to each other, so a naive "paragraphs first,
    tables after" walk (the previous approach here) loses which heading a table
    actually sits under — every table ends up generically labelled "Table N" with
    no connection to a heading like "Pricing & Timeline Estimates". Walking the
    underlying body XML in document order lets each table inherit the heading
    that precedes it, which both the embedding and the LLM need to know a table
    is "the pricing table" rather than an unlabelled grid of numbers.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _docx_page_count(file_path: str, word_count: int) -> int:
    """Word/LibreOffice stamp the actual rendered page count into the docx's
    docProps/app.xml <Pages> element every time the file is saved — read that
    first since it's exact. Fall back to a ~500-words-per-page estimate (the
    standard single-spaced Word default) only when that property is missing,
    e.g. a docx that was generated programmatically and never opened/saved in
    a word processor — an estimate beats the previous behaviour of always
    showing 0 pages regardless of document length."""
    import zipfile
    from xml.etree import ElementTree as ET

    try:
        with zipfile.ZipFile(file_path) as z:
            with z.open("docProps/app.xml") as f:
                ns = {"ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"}
                pages_el = ET.parse(f).find("ep:Pages", ns)
                if pages_el is not None and (pages_el.text or "").isdigit():
                    pages = int(pages_el.text)
                    if pages > 0:
                        return pages
    except Exception:
        pass
    return max(1, round(word_count / 500)) if word_count else 0


def _extract_docx(file_path: str, document: Document) -> List[dict]:
    """Extract paragraphs and tables from a DOCX file in document order, tagging
    each with the most recent heading seen (its section_title)."""
    from docx import Document as DocxDocument
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    doc = DocxDocument(file_path)
    pages_data = []
    current_section: Optional[str] = None
    buffer: list[str] = []
    table_count = 0

    def flush_buffer():
        if buffer:
            pages_data.append({
                "text": " ".join(buffer),
                "page_number": None,
                "section_title": current_section,
            })
            buffer.clear()

    for block in _iter_block_items(doc):
        if isinstance(block, DocxParagraph):
            text = block.text.strip()
            if not text:
                continue
            style_name = block.style.name.lower() if block.style else ""
            if "heading" in style_name:
                flush_buffer()
                current_section = text
            else:
                buffer.append(text)
        elif isinstance(block, DocxTable):
            table_count += 1
            rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
            if not _is_meaningful_table(rows):
                continue
            # A table interrupts the running paragraph buffer — flush what came
            # before it under the current heading, then resume buffering after.
            flush_buffer()
            heading = current_section or f"Table {table_count}"
            html = _table_to_html(rows)
            pages_data.append({
                # Prefixing the heading into the embedded text (not just the
                # metadata) is what actually helps retrieval: an embedding of
                # bare "<table><tr><td>Components</td>..." carries none of the
                # "this is pricing" signal a query like "costing" needs, while
                # "Pricing & Timeline Estimates:\n<table>..." does.
                "text": f"{heading}:\n{html}" if current_section else html,
                "page_number": None,
                "section_title": heading,
                "is_table": True,
            })

    flush_buffer()
    word_count = sum(len(p["text"].split()) for p in pages_data if not p.get("is_table"))
    document.page_count = _docx_page_count(file_path, word_count)
    return pages_data


def _extract_xlsx(file_path: str, document: Document) -> List[dict]:
    """Extract each sheet of an XLSX file as one atomic HTML table chunk."""
    import openpyxl

    wb = openpyxl.load_workbook(file_path, data_only=True)
    document.page_count = len(wb.sheetnames)  # "pages" ~= sheets for a spreadsheet
    pages_data = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [
            [str(cell) if cell is not None else "" for cell in row]
            for row in ws.iter_rows(values_only=True)
            if any(cell is not None and str(cell).strip() for cell in row)
        ]
        if not rows:
            continue
        pages_data.append({
            "text": _table_to_html(rows),
            "page_number": None,
            "section_title": sheet_name,
            "is_table": True,
        })
    return pages_data


def _extract_pptx(file_path: str, document: Document) -> List[dict]:
    """Extract text from each slide of a PPTX file."""
    from pptx import Presentation

    prs = Presentation(file_path)
    document.page_count = len(prs.slides)  # "pages" ~= slides for a deck
    pages_data = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = " ".join(run.text for run in para.runs).strip()
                    if line:
                        texts.append(line)
        if texts:
            pages_data.append({
                "text": " ".join(texts),
                "page_number": slide_num,
                "section_title": f"Slide {slide_num}",
            })
    return pages_data


# ── Main ingestion function ───────────────────────────────────────────────────

async def ingest_document(file_path: str, document_id: str) -> None:
    """
    Full ingestion pipeline for a single document.

    Opens its own DB session so it is not affected by the request-scoped
    session closing when the upload response is returned.

    Steps:
    1. Detect file type and extract text
    2. Chunk extracted text
    3. Embed chunks with SentenceTransformer
    4. Upsert to Pinecone
    5. Save Chunk records to SQLite
    6. Mark document.is_indexed = True
    """
    from app.database import SessionLocal

    db = SessionLocal()
    try:
        document = db.query(Document).filter(Document.id == document_id).first()
        if not document:
            logger.error("ingest_document: document %s not found in DB", document_id)
            return

        logger.info("Starting ingestion for document %s (%s)", document.id, document.file_type)
        document.ingestion_error = None  # clear any error from a previous attempt

        # Step 1: Extract text (legacy .doc/.ppt/.xls are converted to OOXML first)
        ft = (document.file_type or "").lower().lstrip(".")
        pages_data: List[dict] = []

        try:
            extract_path = file_path
            if is_legacy_office_format(ft):
                extract_path = convert_legacy_office_doc(file_path, ft)

            if ft == "pdf":
                pages_data = _extract_pdf(extract_path, document)
            elif ft in ("docx", "doc"):
                pages_data = _extract_docx(extract_path, document)
            elif ft in ("xlsx", "xls"):
                pages_data = _extract_xlsx(extract_path, document)
            elif ft in ("pptx", "ppt"):
                pages_data = _extract_pptx(extract_path, document)
            else:
                with open(extract_path, "r", errors="replace") as f:
                    raw = f.read()
                pages_data = [{"text": raw, "page_number": None, "section_title": None}]
        except Exception as exc:
            logger.error("Text extraction failed for document %s: %s", document.id, exc, exc_info=True)
            document.ingestion_error = str(exc)[:1000]
            db.commit()
            return

        # Step 1b: AI classification — suggest a category from the extracted content
        # so the uploader isn't required to pick one manually. Skipped if a human
        # already confirmed a category (e.g. re-ingesting a new version).
        if not document.category_confirmed:
            sample_text = "\n\n".join(p["text"] for p in pages_data[:3] if p.get("text"))
            from app.documents.classifier import classify_document
            classification = classify_document(sample_text, document.file_name)
            document.doc_type = classification["category"]
            document.ai_detected_type = classification["detected_type"]
            db.commit()
            logger.info(
                "Document %s classified as %s (%s)",
                document.id, classification["category"], classification["detected_type"],
            )

        # Step 2: Chunk
        all_chunks: List[dict] = []
        for page in pages_data:
            if page.get("is_table"):
                # Keep the whole table as one chunk — running it through the
                # word-count chunker would slice rows apart at arbitrary points.
                words = page["text"].split()
                if words:
                    all_chunks.append({
                        "text": page["text"],
                        "chunk_index": 0,
                        "page_number": page.get("page_number"),
                        "section_title": page.get("section_title"),
                        "token_count": len(words),
                        "is_table": True,
                    })
                continue
            for chunk in chunk_text(
                page["text"],
                page_number=page.get("page_number"),
                section_title=page.get("section_title"),
            ):
                all_chunks.append(chunk)

        if not all_chunks:
            logger.warning("No chunks extracted for document %s", document.id)
            document.ingestion_error = (
                "No extractable text was found in this document — it may be empty, "
                "image-only, or corrupted."
            )
            db.commit()
            db.close()
            return

        # Step 3: Embed
        texts = [c["text"] for c in all_chunks]
        embeddings = embed_texts(texts)

        # Step 4 & 5: Build Pinecone records and SQLite Chunk rows
        upload_date = document.created_at.isoformat() if document.created_at else datetime.now(timezone.utc).isoformat()
        pinecone_vectors: List[dict] = []
        db_chunks: List[Chunk] = []

        for i, (chunk, embedding) in enumerate(zip(all_chunks, embeddings)):
            chunk_id = str(uuid.uuid4())

            # Pinecone record
            vector_record = {
                "id": chunk_id,
                "values": embedding,
                "metadata": {
                    "document_id": document.id,
                    "owner_user_id": str(document.uploader_id),
                    "division": document.division or "",
                    "project": document.project or "",
                    "confidentiality": document.confidentiality,
                    "doc_type": document.doc_type or "",
                    "is_hr_data": False,
                    "chunk_index": chunk["chunk_index"],
                    "page_number": chunk.get("page_number") or 0,
                    "section_title": chunk.get("section_title") or "",
                    "text_preview": chunk["text"][:200],
                    "file_type": ft,
                    "is_scanned": document.is_scanned,
                    "ocr_confidence": document.ocr_confidence or 0.0,
                    "upload_date": upload_date,
                    "version": document.current_version,
                    "document_title": document.title,
                    # Lets retrieval give tabular chunks a relevance boost — pricing,
                    # deliverable, and spec tables are exactly the content most likely
                    # to hold the precise figure a query is after, yet a table's dense
                    # HTML/markup embeds and keyword-matches worse than plain prose
                    # saying the same thing in sentences, so it's otherwise structurally
                    # under-ranked against boilerplate that just happens to repeat the
                    # query's words more.
                    "is_table": bool(chunk.get("is_table", False)),
                },
            }
            pinecone_vectors.append(vector_record)

            # SQLite Chunk record
            db_chunks.append(
                Chunk(
                    id=chunk_id,
                    document_id=document.id,
                    chunk_index=chunk["chunk_index"],
                    text=chunk["text"],
                    page_number=chunk.get("page_number"),
                    section_title=chunk.get("section_title"),
                    token_count=chunk["token_count"],
                )
            )

        # Upsert to Pinecone
        pinecone_client.upsert_chunks(pinecone_vectors)

        # Save chunks to DB
        try:
            db.add_all(db_chunks)
            db.commit()
        except Exception as exc:
            logger.error("Failed to save chunks to DB for doc %s: %s", document.id, exc)
            db.rollback()

        # Step 6: Mark indexed
        document.is_indexed = True
        db.commit()
        logger.info(
            "Ingestion complete for document %s: %d chunks indexed",
            document.id,
            len(all_chunks),
        )

    except Exception as exc:
        logger.error("Ingestion failed for document %s: %s", document_id, exc, exc_info=True)
        db.rollback()
        # Best-effort: surface the failure on the document rather than leaving it stuck
        # at "pending" forever with no explanation.
        try:
            doc = db.query(Document).filter(Document.id == document_id).first()
            if doc:
                doc.ingestion_error = str(exc)[:1000]
                db.commit()
        except Exception:
            db.rollback()
    finally:
        db.close()
